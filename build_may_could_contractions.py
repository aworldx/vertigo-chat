from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = "/Users/amirhasanov/Projects/chat/english_may_could_contractions.docx"


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:color"), "D9D9D9")


def set_margins(cell, top=100, start=135, bottom=100, end=135):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_text(p, text, bold=False, size=10, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    return r


def put_cell(cell, lines, header=False, center=False):
    p = cell.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.04
    for i, (text, bold) in enumerate(lines):
        if i:
            p.add_run("\n")
        add_text(p, text, bold=bold, size=9.4 if not header else 9.7, color=(255, 255, 255) if header else None)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_margins(cell)
    set_border(cell)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    add_text(p, text, bold=True, size=12)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.25)
section.bottom_margin = Cm(1.25)
section.left_margin = Cm(1.35)
section.right_margin = Cm(1.35)

doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
title_pr = doc.styles["Title"]._element.get_or_add_pPr()
title_border = title_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_pr.remove(title_border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(3)
add_text(title, "May Could и сокращения", bold=True, size=20, color=(0, 0, 0))

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.paragraph_format.space_after = Pt(8)
add_text(intro, "Вежливые просьбы и короткая запись is и has", bold=True, size=10.5)

add_heading(doc, "May и Could для вежливой просьбы")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
p.paragraph_format.line_spacing = 1.08
add_text(p, "may", bold=True, size=10.5)
add_text(p, " [мэй] и ", size=10.5)
add_text(p, "could", bold=True, size=10.5)
add_text(p, " [куд] помогают вежливо о чём-то попросить. Обычно ", size=10.5)
add_text(p, "May I...?", bold=True, size=10.5)
add_text(p, " звучит очень вежливо и официально, а ", size=10.5)
add_text(p, "Could I...?", bold=True, size=10.5)
add_text(p, " или ", size=10.5)
add_text(p, "Could you...?", bold=True, size=10.5)
add_text(p, " - вежливо и обычно в разговоре.", size=10.5)

may_table = doc.add_table(rows=1, cols=3)
may_table.autofit = False
may_table.style = "Table Grid"
may_widths = [Cm(4.6), Cm(8.4), Cm(5.1)]
for cell, width, label in zip(may_table.rows[0].cells, may_widths, ["Структура", "Пример", "Перевод"]):
    cell.width = width
    set_shading(cell, "244C66")
    put_cell(cell, [(label, True)], header=True, center=True)
may_rows = [
    ([("May I + действие?", True)], [("May I come in?", True), ("[мэй ай кам ин?]", False)], [("Можно войти?", False)]),
    ([("Could I + действие?", True)], [("Could I have a pencil?", True), ("[куд ай хэв э пэнсл?]", False)], [("Можно мне карандаш?", False)]),
    ([("Could you + действие?", True)], [("Could you help me?", True), ("[куд ю: хелп ми?]", False)], [("Ты можешь мне помочь?", False)]),
    ([("May I + действие?", True)], [("May I open the window?", True), ("[мэй ай оупэн зэ уиндоу?]", False)], [("Можно открыть окно?", False)]),
]
for index, row_data in enumerate(may_rows):
    cells = may_table.add_row().cells
    for cell, width, content in zip(cells, may_widths, row_data):
        cell.width = width
        if index % 2:
            set_shading(cell, "EFF5F8")
        put_cell(cell, content, center=(cell == cells[0]))

add_heading(doc, "Сокращение с буквой s")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
p.paragraph_format.line_spacing = 1.08
add_text(p, "Апостроф ", size=10.5)
add_text(p, "'", bold=True, size=10.5)
add_text(p, " и буква ", size=10.5)
add_text(p, "s", bold=True, size=10.5)
add_text(p, " могут скрывать короткое слово. Чаще всего ", size=10.5)
add_text(p, "'s", bold=True, size=10.5)
add_text(p, " означает ", size=10.5)
add_text(p, "is", bold=True, size=10.5)
add_text(p, " [из] или ", size=10.5)
add_text(p, "has", bold=True, size=10.5)
add_text(p, " [хэз].", size=10.5)

short_table = doc.add_table(rows=1, cols=4)
short_table.autofit = False
short_table.style = "Table Grid"
short_widths = [Cm(4.75), Cm(4.75), Cm(4.8), Cm(3.8)]
for cell, width, label in zip(short_table.rows[0].cells, short_widths, ["Полная форма", "Коротко", "Как читаем", "Что скрыто"]):
    cell.width = width
    set_shading(cell, "244C66")
    put_cell(cell, [(label, True)], header=True, center=True)
short_rows = [
    ([("She is happy.", True)], [("She's happy.", True)], [("[ши:з хэпи]", False)], [("she is", True), ("она счастливая", False)]),
    ([("He is a doctor.", True)], [("He's a doctor.", True)], [("[хи:з э доктэ]", False)], [("he is", True), ("он врач", False)]),
    ([("It is a cat.", True)], [("It's a cat.", True)], [("[итс э кэт]", False)], [("it is", True), ("это кошка", False)]),
    ([("She has got a dress.", True)], [("She's got a dress.", True)], [("[ши:з гот э дрэс]", False)], [("she has", True), ("у неё есть платье", False)]),
    ([("Your mom has got a car.", True)], [("Your mom's got a car.", True)], [("[йо: мамз гот э ка:]", False)], [("mom has", True), ("у мамы есть машина", False)]),
]
for index, row_data in enumerate(short_rows):
    cells = short_table.add_row().cells
    for cell, width, content in zip(cells, short_widths, row_data):
        cell.width = width
        if index % 2:
            set_shading(cell, "EFF5F8")
        put_cell(cell, content, center=(cell == cells[2]))

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(2)
p.paragraph_format.line_spacing = 1.06
add_text(p, "Как понять: ", bold=True, size=10)
add_text(p, "если после сокращения идёт ", size=10)
add_text(p, "got", bold=True, size=10)
add_text(p, ", то обычно скрыто ", size=10)
add_text(p, "has", bold=True, size=10)
add_text(p, ". Если дальше признак или кто-то, то обычно скрыто ", size=10)
add_text(p, "is", bold=True, size=10)
add_text(p, ".", size=10)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.06
add_text(p, "Важно: ", bold=True, size=10)
add_text(p, "в выражении ", size=10)
add_text(p, "Mom's car", bold=True, size=10)
add_text(p, " [мамз ка:] слово ", size=10)
add_text(p, "'s", bold=True, size=10)
add_text(p, " означает «мамин»: мамина машина. Здесь это не ", size=10)
add_text(p, "is", bold=True, size=10)
add_text(p, " и не ", size=10)
add_text(p, "has", bold=True, size=10)
add_text(p, ".", size=10)

doc.save(OUT)
print(OUT)
