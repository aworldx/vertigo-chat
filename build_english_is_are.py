from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = "/Users/amirhasanov/Projects/chat/english_is_are_cheatsheet.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def clear_cell(cell):
    cell.text = ""
    return cell.paragraphs[0]


def add_text(paragraph, text, bold=False, size=10.2, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def add_cell_text(cell, lines, header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = clear_cell(cell)
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.04
    for i, (text, bold) in enumerate(lines):
        if i:
            p.add_run("\n")
        add_text(p, text, bold=bold, size=9.3 if not header else 9.5, color=(255, 255, 255) if header else None)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.25)
section.bottom_margin = Cm(1.25)
section.left_margin = Cm(1.35)
section.right_margin = Cm(1.35)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
styles["Normal"].font.size = Pt(10.5)
title_style_pr = styles["Title"]._element.get_or_add_pPr()
title_style_border = title_style_pr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_pr.remove(title_style_border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)
run = title.add_run("IS ARE HAVE GOT")
run.font.name = "Arial"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
run.font.size = Pt(20)
run.font.bold = True
run.font.color.rgb = RGBColor(0, 0, 0)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(10)
add_text(subtitle, "Памятка по утвердительным и вопросительным предложениям", bold=True, size=10.5)

heading = doc.add_paragraph()
heading.paragraph_format.space_before = Pt(0)
heading.paragraph_format.space_after = Pt(3)
add_text(heading, "1. Зачем нужны is и are", bold=True, size=12)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
p.paragraph_format.line_spacing = 1.08
add_text(p, "В английском предложении обычно нужен глагол. ", size=10.5)
add_text(p, "is", bold=True, size=10.5)
add_text(p, " [из] и ", size=10.5)
add_text(p, "are", bold=True, size=10.5)
add_text(p, " [а:] соединяют того, о ком говорят, с информацией о нём: кто он, какой он или где он. В русском языке это слово часто не произносится.", size=10.5)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
p.paragraph_format.line_spacing = 1.08
add_text(p, "They are doctors.", bold=True, size=10.5)
add_text(p, " [зэй а: доктэз] - Они врачи.  Нельзя: ", size=10.5)
add_text(p, "They doctors.", bold=True, size=10.5)

heading = doc.add_paragraph()
heading.paragraph_format.space_before = Pt(3)
heading.paragraph_format.space_after = Pt(3)
add_text(heading, "2. Как работает have got", bold=True, size=12)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
p.paragraph_format.line_spacing = 1.08
add_text(p, "have got", bold=True, size=10.5)
add_text(p, " [хэв гот] и ", size=10.5)
add_text(p, "has got", bold=True, size=10.5)
add_text(p, " [хэз гот] - это способ сказать «у кого-то есть». В вопросе вперёд переходит только ", size=10.5)
add_text(p, "have", bold=True, size=10.5)
add_text(p, " [хэв] или ", size=10.5)
add_text(p, "has", bold=True, size=10.5)
add_text(p, ", а ", size=10.5)
add_text(p, "got", bold=True, size=10.5)
add_text(p, " [гот] остаётся после человека.", size=10.5)

heading = doc.add_paragraph()
heading.paragraph_format.space_before = Pt(3)
heading.paragraph_format.space_after = Pt(3)
add_text(heading, "3. Порядок слов", bold=True, size=12)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
p.paragraph_format.line_spacing = 1.08
add_text(p, "Утверждение: ", bold=True, size=10.5)
add_text(p, "кто? + нужное слово + информация. ", size=10.5)
add_text(p, "They are doctors.", bold=True, size=10.5)
add_text(p, " [зэй а: доктэз] или ", size=10.5)
add_text(p, "I have got a dress.", bold=True, size=10.5)
add_text(p, " [ай хэв гот э дрэс].", size=10.5)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
p.paragraph_format.line_spacing = 1.08
add_text(p, "Вопрос: ", bold=True, size=10.5)
add_text(p, "is [из], are [а:], have [хэв] или has [хэз] ставим на первое место. ", size=10.5)
add_text(p, "Are they doctors?", bold=True, size=10.5)
add_text(p, " [а: зэй доктэз?]  ", size=10.5)
add_text(p, "Have you got a pencil?", bold=True, size=10.5)
add_text(p, " [хэв ю: гот э пэнсл?]", size=10.5)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
p.paragraph_format.line_spacing = 1.08
add_text(p, "Короткий ответ: ", bold=True, size=10.5)
add_text(p, "повторяем нужное слово из вопроса: ", size=10.5)
add_text(p, "Are they...? - Yes, they are.", bold=True, size=10.5)
add_text(p, " [йес, зэй а:].  ", size=10.5)
add_text(p, "Have you...? - No, I haven't.", bold=True, size=10.5)
add_text(p, " [ноу, ай хэвнт].", size=10.5)

heading = doc.add_paragraph()
heading.paragraph_format.space_before = Pt(3)
heading.paragraph_format.space_after = Pt(5)
add_text(heading, "Сводная таблица", bold=True, size=12)

table = doc.add_table(rows=1, cols=4)
table.autofit = False
table.style = "Table Grid"
widths = [Cm(3.05), Cm(4.55), Cm(5.2), Cm(4.5)]
for cell, width in zip(table.rows[0].cells, widths):
    cell.width = width
headers = ["Когда", "Утверждение", "Вопрос", "Короткий ответ"]
for cell, text in zip(table.rows[0].cells, headers):
    set_cell_shading(cell, "244C66")
    add_cell_text(cell, [(text, True)], header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
set_repeat_table_header(table.rows[0])

rows = [
    (
        [("Один человек или предмет", True), ("he [хи:], she [ши:], it [ит]", False)],
        [("She is a doctor.", True), ("[ши: из э доктэ]", False), ("Она врач.", False)],
        [("Is she a doctor?", True), ("[из ши: э доктэ?]", False)],
        [("Yes, she is.", True), ("[йес, ши: из]", False), ("No, she isn’t.", True), ("[ноу, ши: изнт]", False)],
    ),
    (
        [("Несколько людей; также you [ю:]", True), ("we [уи:], they [зэй]", False)],
        [("They are doctors.", True), ("[зэй а: доктэз]", False), ("Они врачи.", False)],
        [("Are they doctors?", True), ("[а: зэй доктэз?]", False)],
        [("Yes, they are.", True), ("[йес, зэй а:]", False), ("No, they aren’t.", True), ("[ноу, зэй а:нт]", False)],
    ),
    (
        [("У меня, тебя, нас, их что-то есть", True), ("I, you, we, they", False)],
        [("I have got a dress.", True), ("[ай хэв гот э дрэс]", False), ("У меня есть платье.", False)],
        [("Have you got a pencil?", True), ("[хэв ю: гот э пэнсл?]", False)],
        [("Yes, I have.", True), ("[йес, ай хэв]", False), ("No, I haven’t.", True), ("[ноу, ай хэвнт]", False)],
    ),
    (
        [("У него, неё или предмета что-то есть", True), ("he, she, it", False)],
        [("Your mom has got a car.", True), ("[йо: мам хэз гот э ка:]", False), ("У твоей мамы есть машина.", False)],
        [("Has your mom got a car?", True), ("[хэз йо: мам гот э ка:?]", False)],
        [("Yes, she has.", True), ("[йес, ши: хэз]", False), ("No, she hasn’t.", True), ("[ноу, ши: хэзнт]", False)],
    ),
]

for index, row_data in enumerate(rows):
    cells = table.add_row().cells
    for cell, width, content in zip(cells, widths, row_data):
        cell.width = width
        if index % 2:
            set_cell_shading(cell, "EFF5F8")
        add_cell_text(cell, content, align=WD_ALIGN_PARAGRAPH.LEFT)

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(6)
note.paragraph_format.space_after = Pt(0)
note.paragraph_format.line_spacing = 1.0
add_text(note, "Важно: ", bold=True, size=9.5)
add_text(note, "в вопросе с have got не говорим “Have got you...?” Правильно: “Have you got...?”", size=9.5)

doc.save(OUT)
print(OUT)
