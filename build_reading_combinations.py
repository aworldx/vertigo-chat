from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = "/Users/amirhasanov/Projects/chat/english_reading_combinations.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=115, bottom=80, end=115):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_run(paragraph, text, bold=False, size=9.5, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def fill_cell(cell, content, header=False, centered=False):
    p = cell.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.03
    for idx, (text, bold) in enumerate(content):
        if idx:
            p.add_run("\n")
        add_run(p, text, bold=bold, size=9.2 if not header else 9.7, color=(255, 255, 255) if header else None)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = Cm(29.7), Cm(21)
section.top_margin = Cm(1.1)
section.bottom_margin = Cm(1.1)
section.left_margin = Cm(1.15)
section.right_margin = Cm(1.15)

doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

title_style_pr = doc.styles["Title"]._element.get_or_add_pPr()
title_style_border = title_style_pr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_pr.remove(title_style_border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(3)
add_run(title, "Сочетания букв в английском", bold=True, size=20, color=(0, 0, 0))

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.paragraph_format.space_after = Pt(8)
add_run(intro, "Читаем сочетание целиком, затем слово. Русская подсказка помогает в начале, но полезно слушать и повторять слово вслух.", bold=True, size=10)

table = doc.add_table(rows=1, cols=4)
table.autofit = False
table.style = "Table Grid"
widths = [Cm(2.9), Cm(5.35), Cm(13.2), Cm(5.95)]
headers = ["Сочетание", "Как читаем", "Примеры", "Значение слов"]
for cell, width, label in zip(table.rows[0].cells, widths, headers):
    cell.width = width
    set_cell_shading(cell, "244C66")
    fill_cell(cell, [(label, True)], header=True, centered=True)

rows = [
    ([("sh", True)], [("[ш]", True)], [("she [ши:]", True), ("fish [фиш]", True), ("shop [шоп]", True)], [("она", False), ("рыба", False), ("магазин", False)]),
    ([("ch", True)], [("[ч]", True)], [("chair [чэа]", True), ("cheese [чи:з]", True), ("chicken [чикэн]", True)], [("стул", False), ("сыр", False), ("курица", False)]),
    ([("th", True)], [("[с]", True), ("кончик языка между зубами", False)], [("three [сри:]", True), ("thin [син]", True), ("thank [сэнк]", True)], [("три", False), ("тонкий", False), ("спасибо", False)]),
    ([("th", True)], [("[з]", True), ("кончик языка между зубами", False)], [("the [зэ]", True), ("this [зис]", True), ("they [зэй]", True)], [("этот, эта, эти", False), ("это", False), ("они", False)]),
    ([("ee", True)], [("[и:]", True)], [("green [гри:н]", True), ("see [си:]", True), ("feet [фи:т]", True)], [("зелёный", False), ("видеть", False), ("ноги, ступни", False)]),
    ([("oo", True)], [("[у:]", True)], [("moon [мун]", True), ("room [ру:м]", True), ("food [фу:д]", True)], [("луна", False), ("комната", False), ("еда", False)]),
    ([("ai", True)], [("[эй]", True)], [("rain [рэйн]", True), ("train [трэйн]", True), ("tail [тэйл]", True)], [("дождь", False), ("поезд", False), ("хвост", False)]),
    ([("ay", True)], [("[эй]", True)], [("play [плэй]", True), ("day [дэй]", True), ("gray [грэй]", True)], [("играть", False), ("день", False), ("серый", False)]),
    ([("oa", True)], [("[оу]", True)], [("boat [боут]", True), ("road [роуд]", True), ("coat [коут]", True)], [("лодка", False), ("дорога", False), ("пальто", False)]),
    ([("ar", True)], [("[а:]", True)], [("car [ка:]", True), ("star [ста:]", True), ("park [па:к]", True)], [("машина", False), ("звезда", False), ("парк", False)]),
]

for index, row_data in enumerate(rows):
    cells = table.add_row().cells
    for cell, width, content in zip(cells, widths, row_data):
        cell.width = width
        if index % 2:
            set_cell_shading(cell, "EFF5F8")
        fill_cell(cell, content, centered=(cell == cells[0] or cell == cells[1]))

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(6)
note.paragraph_format.space_after = Pt(0)
note.paragraph_format.line_spacing = 1.0
add_run(note, "Важно: ", bold=True, size=9.3)
add_run(note, "сочетания обычно читаются так, но в английском бывают исключения. Например, oo в book [бук] читается не так, как в moon [мун].", size=9.3)

doc.save(OUT)
print(OUT)
